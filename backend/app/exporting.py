"""이벤트 내보내기 파일 만들기 (CSV · 엑셀 · 스냅샷 묶음)"""

import csv
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 넣을 수 있는 항목 (화면의 체크 상자와 같은 순서)
COLUMNS: dict[str, str] = {
    "id": "번호",
    "time": "시각",
    "camera": "카메라",
    "zone": "구역",
    "type": "종류",
    "severity": "심각도",
    "confidence": "신뢰도",
    "status": "처리 상태",
    "handler": "처리 담당자",
    "memo": "조치 메모",
    "snapshot": "스냅샷 경로",
}
DEFAULT_COLUMNS = ["id", "time", "camera", "zone", "type", "severity", "confidence", "status"]
WIDTHS = {"id": 10, "time": 20, "camera": 12, "zone": 22, "type": 16, "severity": 8, "confidence": 8,
          "status": 10, "handler": 12, "memo": 40, "snapshot": 40}
SEVERITY_FILL = {"높음": "FFE3E3", "중간": "FFF1D6", "낮음": "E3F6EA"}


def to_csv(columns: list[str], rows: list[dict]) -> bytes:
    """엑셀에서 한글이 깨지지 않게 BOM 을 붙인 UTF-8"""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([COLUMNS[c] for c in columns])
    for r in rows:
        w.writerow([r[c] for c in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(columns: list[str], rows: list[dict], title: str) -> bytes:
    """보고용 엑셀: 머리글 고정 · 자동 필터 · 심각도 칸 색"""
    wb = Workbook()
    ws = wb.active
    ws.title = "이벤트"
    ws.append([COLUMNS[c] for c in columns])
    for r in rows:
        ws.append([r[c] for c in columns])
    head_fill = PatternFill("solid", fgColor="1C2430")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 12)
    if "severity" in columns:
        col = columns.index("severity") + 1
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            fill = SEVERITY_FILL.get(row[0].value)
            if fill:
                row[0].fill = PatternFill("solid", fgColor=fill)
    if "memo" in columns:
        col = columns.index("memo") + 1
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.properties.title = title
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_zip(columns: list[str], rows: list[dict], snapshot_files: dict[int, Path]) -> bytes:
    """CSV + 이벤트별 탐지 이미지. CSV 의 스냅샷 경로는 묶음 안의 파일 이름으로 바꾼다"""
    out = io.BytesIO()
    cols = columns if "snapshot" in columns else [*columns, "snapshot"]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        packed = []
        for r in rows:
            path = snapshot_files.get(r["_event_id"])
            if path and path.is_file():
                name = f"snapshots/EV-{r['_event_id']}{path.suffix}"
                z.write(path, name)
                packed.append({**r, "snapshot": name})
            else:
                packed.append({**r, "snapshot": "-"})
        z.writestr("events.csv", to_csv(cols, packed))
    return out.getvalue()
