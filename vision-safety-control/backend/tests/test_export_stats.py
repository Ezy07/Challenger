import csv
import io
import zipfile
from datetime import date, timedelta

from openpyxl import load_workbook

TODAY = date.today()


# ---------------------------------------------------------------- 내보내기
def test_export_preview(client, manager):
    res = client.get("/api/events/export/preview", params=[("columns", "id"), ("columns", "type"), ("columns", "memo")],
                     headers=manager).json()
    assert res["total"] == 9
    assert [c["label"] for c in res["columns"]] == ["번호", "종류", "조치 메모"]
    assert len(res["rows"]) == 5
    assert all(r[0].startswith("EV-") for r in res["rows"])
    bad = client.get("/api/events/export/preview", params={"columns": "password"}, headers=manager)
    assert bad.status_code == 422


def test_export_csv_with_handler_and_memo(client, manager):
    params = [("format", "csv"), ("severity", "HIGH"), ("columns", "id"), ("columns", "handler"), ("columns", "memo")]
    res = client.get("/api/events/export", params=params, headers=manager)
    assert res.status_code == 200 and "attachment" in res.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))
    assert rows[0] == ["번호", "처리 담당자", "조치 메모"]
    assert len(rows) == 1 + 3
    done = [r for r in rows[1:] if r[2] != "-"]
    assert done and done[0][1] == "오세영"  # 처리완료된 장비 근접 이벤트


def test_export_xlsx(client, manager):
    res = client.get("/api/events/export", params={"format": "xlsx"}, headers=manager)
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in ws[1]] == ["번호", "시각", "카메라", "구역", "종류", "심각도", "신뢰도", "처리 상태"]
    assert ws.max_row == 1 + 9
    assert ws.freeze_panes == "A2"


def test_export_zip_contains_snapshots(client, manager, tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "snapshot_dir", tmp_path / "snapshots")  # seed_demo 가 예시 스냅샷을 만든 곳
    res = client.get("/api/events/export", params={"format": "zip"}, headers=manager)
    assert res.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(res.content))
    names = z.namelist()
    assert "events.csv" in names
    images = [n for n in names if n.startswith("snapshots/")]
    assert len(images) == 8  # 카메라가 있는 이벤트 8건 (기상 경보 1건은 카메라 없음)
    rows = list(csv.DictReader(io.StringIO(z.read("events.csv").decode("utf-8-sig"))))
    assert {r["스냅샷 경로"] for r in rows} - {"-"} == set(images)


def test_snapshot_path_cannot_escape_folder():
    from app.routers.events import snapshot_file

    assert snapshot_file("snapshots/../../.env") is None
    assert snapshot_file("/etc/passwd") is None


def test_event_list_has_snapshot(client, manager):
    items = client.get("/api/events", headers=manager).json()["items"]
    assert items[0]["snapshot"].startswith("snapshots/")


# ---------------------------------------------------------------- 통계
def test_summary_today(client, manager):
    s = client.get("/api/stats/summary", headers=manager).json()
    cur = s["current"]
    assert (cur["total"], cur["high"], cur["open"], cur["handled_rate"]) == (9, 3, 3, 67)
    assert cur["avg_handle_min"] == 8.8  # 처리완료 5건 10분 · 오탐 1건 3분
    assert s["previous"]["total"] == 0
    assert s["by_zone"][0] == {"key": s["by_zone"][0]["key"], "label": "101동 갱폼 인양구역", "count": 3}
    assert any(z["label"] == "구역 없음" for z in s["by_zone"])  # 기상 경보
    assert s["by_hour"][14] == 3
    assert len(s["by_type"]) == 9 and sum(t["count"] for t in s["by_type"]) == 9


def test_trend_and_report_with_history(client, manager, db_session):
    from app.seed import seed_history

    with db_session() as db:
        made = seed_history(db, 10)
        db.commit()
        assert made > 0 and seed_history(db, 10) == 0  # 두 번 넣지 않음

    t = client.get("/api/stats/trend", params={"days": 14}, headers=manager).json()
    assert len(t["daily"]) == 14 and t["daily"][-1]["key"] == str(TODAY)
    assert t["total"] == made + 9
    assert sum(d["count"] for d in t["daily"]) == t["total"] == sum(t["by_hour"])
    assert t["daily"][0]["count"] == 0  # 14일 전은 기록 없음

    only_high = client.get("/api/stats/trend", params={"days": 7, "severity": "HIGH"}, headers=manager).json()
    assert 0 < only_high["total"] < t["total"]

    # 달력으로 고른 기간 (시작 · 끝 날짜)
    start = TODAY - timedelta(days=4)
    ranged = client.get("/api/stats/trend", params={"date_from": str(start), "date_to": str(TODAY)}, headers=manager).json()
    assert [d["key"] for d in ranged["daily"]] == [str(start + timedelta(days=i)) for i in range(5)]
    assert ranged["avg_per_day"] == round(ranged["total"] / 5, 1)
    bad = client.get("/api/stats/trend", params={"date_from": str(TODAY), "date_to": str(start)}, headers=manager)
    assert bad.status_code == 422  # 끝이 시작보다 빠름

    week_ago = TODAY - timedelta(days=6)
    r = client.get("/api/stats/report", params={"date_from": str(week_ago), "date_to": str(TODAY)}, headers=manager).json()
    assert r["current"]["total"] > 9 and r["previous"]["total"] > 0
    assert len(r["high_events"]) == r["current"]["high"]
    past_high = [h for h in r["high_events"] if h["detected_at"][:10] != str(TODAY)]
    assert past_high and all(h["memo"] for h in past_high)  # 지난 높음 이벤트는 모두 처리 메모가 있음

    bad = client.get("/api/stats/report", params={"date_from": str(TODAY), "date_to": str(week_ago)}, headers=manager)
    assert bad.status_code == 422


def test_history_is_consistent(client, manager, db_session):
    """지난 기록도 처리완료면 해제 항목이 모두 허가돼 있고 RESOLVE 로 끝난다"""
    from app.seed import seed_history

    with db_session() as db:
        seed_history(db, 5)
        db.commit()
    yesterday = str(TODAY - timedelta(days=1))
    items = client.get("/api/events", params={"date_from": yesterday, "date_to": yesterday, "status": "DONE",
                                              "detection_type": "LOAD_OVERHEAD_WORKER"}, headers=manager).json()["items"]
    for row in items:
        detail = client.get(f"/api/events/{row['id']}", headers=manager).json()
        assert all(c["permitted"] for c in detail["check_items"])
        assert detail["actions"][-1]["action_type"] == "RESOLVE"


def test_export_escapes_formulas():
    from app.exporting import safe_cell, to_csv

    assert safe_cell('=HYPERLINK("http://x","a")') == '\'=HYPERLINK("http://x","a")'
    assert [safe_cell(v) for v in ["+1", "-1+2", "@SUM(A1)", "-", "보통 메모", 0.9, None]] == \
        ["'+1", "'-1+2", "'@SUM(A1)", "-", "보통 메모", 0.9, None]
    rows = list(csv.reader(io.StringIO(to_csv(["memo"], [{"memo": "=1+1"}]).decode("utf-8-sig"))))
    assert rows[1] == ["'=1+1"]
